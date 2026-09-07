#!/usr/bin/env python3
"""
Jyothi Foods — decode the client's two spreadsheets into importer-ready CSVs.

  sections.csv         code, name, sort_order
  items.csv            item_code, pack_type, name, units_per_box, pieces_per_unit,
                       mrp_per_piece, section_code          (price list items)
  unmatched_items.csv  same columns + opening_boxes         (in the stock report but
                       not on the price list; units_per_box only where the name
                       encodes it, e.g. "(12) 32" → 32 — the rest need the client)
  opening_stock.csv    item_code, section_code, opening_boxes

These go through the in-app importer (Setup → Import data), never straight into
the database. This script only regenerates them from docs/reference/.

Usage:
    python3 scripts/import_masters.py docs/reference/Structured_Item_Price_List.xlsx \
        docs/reference/STOCK_REPORT.xlsx seed/
"""
import csv, re, sys
from pathlib import Path
from openpyxl import load_workbook

PACK_MAP = {"JAR": "JAR", "PACK": "PACK", "L.B": "L.B", "LB": "L.B",
            "KG": "KG", "TRY": "TRY", "TRAY": "TRY", "BOX": "BOX"}

# "  8 x 1", "21 X 1", "10 X 1"  ->  8, 21, 10
BOX_RE = re.compile(r"^\s*(\d+)\s*[xX]\s*(\d+)\s*$")
# leading "5/-", "2/-", "10/-"
MRP_RE = re.compile(r"^\s*(\d+)\s*/-")
# the "(12)" packing hint inside the name
PIECES_RE = re.compile(r"\((\d+)\s*[A-Za-z]?\)")
# trailing units-per-box after the parens: "(12) 32", "(12)24 NEW", "(12) 24P"
TRAILING_RE = re.compile(r"\)\s*(\d+)\s*[A-Za-z]?\s*(?:NEW)?\s*$", re.IGNORECASE)


def norm_code(v):
    """'27 A' / '06: A' / '01: A' -> '27A' / '06A' / '01A'. Codes are text, not ints."""
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"[\s:]+", "", s).upper() or None


def clean_name(name):
    return re.sub(r"\s+", " ", str(name).strip())


def parse_name(name):
    """MRP, pieces per unit and (if encoded) units per box from the item name."""
    mrp = MRP_RE.match(name)
    pieces = PIECES_RE.search(name)
    trailing = TRAILING_RE.search(name)
    return (
        float(mrp.group(1)) if mrp else None,
        int(pieces.group(1)) if pieces else 1,
        int(trailing.group(1)) if trailing else None,
    )


def parse_price_list(path):
    """The sheet has two side-by-side blocks: cols A-D and F-I."""
    ws = load_workbook(path, read_only=True, data_only=True)["Item List"]
    items, seen = [], set()

    for row in ws.iter_rows(values_only=True):
        row = list(row) + [None] * 10
        for base in (0, 5):                       # two blocks per row
            code, pack, name, box = row[base:base + 4]
            code, name = norm_code(code), (clean_name(name) if name else None)
            if not code or not name or code == "CODE":
                continue
            if code in seen:                       # first occurrence wins
                continue

            m = BOX_RE.match(str(box or ""))
            if not m:
                continue
            units_per_box = int(m.group(1))
            pack_type = PACK_MAP.get(str(pack or "").strip().upper(), "JAR")
            mrp, pieces, _ = parse_name(name)

            seen.add(code)
            items.append({
                "item_code": code,
                "pack_type": pack_type,
                "name": name,
                "units_per_box": units_per_box,
                "pieces_per_unit": pieces,
                "mrp_per_piece": mrp,
                "section_code": "",
            })
    return items


def parse_stock_report(path):
    """Rows are grouped under section headers like 'S-10 RAMA KRISHNA MESTRI'."""
    ws = load_workbook(path, read_only=True, data_only=True)["Sheet1"]
    sections, stock = [], []
    cur_code = cur_name = None
    order = 0

    for row in ws.iter_rows(values_only=True):
        row = list(row) + [None] * 8
        # A section header has NO item code and NO pack — only the label cell.
        # (An item row with a blank Opening is still an item row.)
        is_header = row[0] is None and row[1] is None and row[2] is not None
        if is_header:
            label = str(row[2]).strip()
            if label.lower().startswith("stock report"):
                continue
            m = re.match(r"^(S-\s*\d+)\s+(.*)$", label)
            if m:
                cur_code, cur_name = m.group(1).replace(" ", ""), m.group(2).strip()
            else:
                cur_code, cur_name = None, label
            order += 1
            sections.append({"code": cur_code or "", "name": cur_name, "sort_order": order})
            continue

        code = norm_code(row[0])
        if not code or code == "ITEMCODE":
            continue
        opening = row[3]
        stock.append({
            "item_code": code,
            "pack": str(row[1]).strip() if row[1] else "",
            "name": clean_name(row[2]) if row[2] else "",
            # importer matches a section by code or by name, so OTHERS groups work too
            "section_code": cur_code or cur_name or "",
            "section_name": cur_name or "OTHERS",
            "opening_boxes": round(float(opening), 3) if isinstance(opening, (int, float)) else 0,
        })
    return sections, stock


def main():
    price_path, stock_path, outdir = sys.argv[1], sys.argv[2], Path(sys.argv[3])
    outdir.mkdir(parents=True, exist_ok=True)

    items = parse_price_list(price_path)
    sections, stock = parse_stock_report(stock_path)

    # attach each item to the section its stock is reported under
    sec_by_item = {}
    for s in stock:
        sec_by_item.setdefault(s["item_code"], s["section_code"])
    for it in items:
        it["section_code"] = sec_by_item.get(it["item_code"], "")

    # items present in the stock report but missing from the price list
    have = {i["item_code"] for i in items}
    orphans, seen = [], set()
    for s in stock:
        if s["item_code"] in have or s["item_code"] in seen:
            continue
        seen.add(s["item_code"])
        mrp, pieces, units = parse_name(s["name"]) if s["name"] else (None, 1, None)
        orphans.append({
            "item_code": s["item_code"],
            "pack_type": "",                      # stock report's Pack column is the stock unit, not the pack type
            "name": s["name"],
            "units_per_box": units if units else "",
            "pieces_per_unit": pieces,
            "mrp_per_piece": mrp,
            "section_code": s["section_code"],
            "opening_boxes": s["opening_boxes"],
        })

    def dump(name, rows, cols):
        with open(outdir / name, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)

    item_cols = ["item_code", "pack_type", "name", "units_per_box",
                 "pieces_per_unit", "mrp_per_piece", "section_code"]
    dump("items.csv", items, item_cols)
    dump("sections.csv", sections, ["code", "name", "sort_order"])
    dump("opening_stock.csv", stock, ["item_code", "section_code", "opening_boxes"])
    dump("unmatched_items.csv", orphans, item_cols + ["opening_boxes"])

    print(f"items             : {len(items)}")
    print(f"sections          : {len(sections)}")
    print(f"opening stock rows: {len(stock)}")
    print(f"in stock report but not in price list: {len(orphans)} "
          f"({sum(1 for o in orphans if o['units_per_box'] != '')} with packing readable from the name)")
    print(f"negative opening  : {sum(1 for s in stock if s['opening_boxes'] < 0)}")


if __name__ == "__main__":
    main()
